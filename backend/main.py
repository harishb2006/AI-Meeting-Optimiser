from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.prompts import ChatPromptTemplate
from langchain.output_parsers import PydanticOutputParser
from dotenv import load_dotenv
import os
import json
from typing import List, Optional
from datetime import datetime
import chromadb
from chromadb.config import Settings

# Load environment variables
load_dotenv()

# Initialize Gemini LLM via LangChain
llm = ChatGoogleGenerativeAI(
    model="gemini-1.5-flash",
    google_api_key=os.getenv("GEMINI_API_KEY"),
    temperature=0.3
)

# Initialize ChromaDB for RAG
chroma_client = chromadb.Client(Settings(
    anonymized_telemetry=False,
    is_persistent=True,
    persist_directory="./chroma_db"
))

# Get or create collection for meetings
try:
    meetings_collection = chroma_client.get_or_create_collection(
        name="meetings",
        metadata={"description": "Meeting transcripts, tasks, and summaries"}
    )
except Exception as e:
    print(f"ChromaDB initialization warning: {e}")
    meetings_collection = None

app = FastAPI(title="AI Meeting Actions Agent")

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://localhost:3001", "http://localhost:3002", "http://localhost:3003"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Pydantic models
class Task(BaseModel):
    task: str
    assignee: Optional[str] = None
    deadline: Optional[str] = None
    priority: Optional[str] = None

class Decision(BaseModel):
    decision: str
    rationale: Optional[str] = None

class TranscriptRequest(BaseModel):
    transcript: str
    meeting_title: str

class AnalysisResponse(BaseModel):
    meeting_title: str
    summary: str
    tasks: List[Task]
    decisions: List[Decision]
    follow_ups: Optional[List[str]] = None

def extract_json_from_response(text: str) -> dict:
    """Extract JSON from markdown code blocks or plain text"""
    try:
        # Try to find JSON in markdown code blocks
        if "```json" in text:
            start = text.find("```json") + 7
            end = text.find("```", start)
            json_text = text[start:end].strip()
        elif "```" in text:
            start = text.find("```") + 3
            end = text.find("```", start)
            json_text = text[start:end].strip()
        else:
            json_text = text.strip()
        
        return json.loads(json_text)
    except json.JSONDecodeError as e:
        # If JSON parsing fails, try to extract it differently
        try:
            # Find the first { and last }
            start = text.find("{")
            end = text.rfind("}") + 1
            if start != -1 and end > start:
                json_text = text[start:end]
                return json.loads(json_text)
        except:
            pass
        raise ValueError(f"Could not parse JSON from response: {str(e)}")

@app.get("/")
def read_root():
    return {"message": "AI Meeting Actions Agent API", "status": "active"}

@app.post("/api/analyze", response_model=AnalysisResponse)
async def analyze_transcript(request: TranscriptRequest):
    """
    Analyze meeting transcript using LangChain and Gemini
    Extract tasks, decisions, summary and store in ChromaDB for RAG
    """
    try:
        # Create the analysis prompt
        prompt_template = """You are an AI assistant that analyzes meeting transcripts and extracts actionable information.

Analyze the following meeting transcript and extract:
1. A concise summary of the meeting (2-3 sentences)
2. All action items/tasks with assignee, deadline (if mentioned), and priority (High/Medium/Low)
3. Key decisions made during the meeting with rationale
4. Any follow-up items or next steps

Meeting Title: {meeting_title}

Transcript:
{transcript}

Provide the response in the following JSON format:
{{
  "summary": "Brief summary of the meeting",
  "tasks": [
    {{
      "task": "Description of the task",
      "assignee": "Person responsible (if mentioned)",
      "deadline": "Deadline (if mentioned)",
      "priority": "High/Medium/Low"
    }}
  ],
  "decisions": [
    {{
      "decision": "Decision made",
      "rationale": "Reason for the decision (if mentioned)"
    }}
  ],
  "follow_ups": [
    "Follow-up item 1",
    "Follow-up item 2"
  ]
}}

Ensure all fields are included even if empty arrays. Be thorough in extracting all tasks and decisions mentioned."""

        # Use LangChain to invoke LLM
        prompt = ChatPromptTemplate.from_template(prompt_template)
        chain = prompt | llm
        
        response = chain.invoke({
            "meeting_title": request.meeting_title,
            "transcript": request.transcript
        })
        
        if not response or not response.content:
            raise HTTPException(status_code=500, detail="No response from Gemini API")
        
        # Extract JSON from response
        analysis_data = extract_json_from_response(response.content)
        
        # Construct the response
        analysis_response = AnalysisResponse(
            meeting_title=request.meeting_title,
            summary=analysis_data.get("summary", "No summary available"),
            tasks=[Task(**task) for task in analysis_data.get("tasks", [])],
            decisions=[Decision(**decision) for decision in analysis_data.get("decisions", [])],
            follow_ups=analysis_data.get("follow_ups", [])
        )
        
        # Store in ChromaDB for RAG (if available)
        if meetings_collection:
            try:
                meeting_id = f"{request.meeting_title}_{datetime.now().isoformat()}"
                
                # Prepare document for storage
                document_text = f"""
Meeting: {request.meeting_title}
Date: {datetime.now().strftime('%Y-%m-%d')}
Summary: {analysis_response.summary}

Tasks:
{chr(10).join([f"- {t.task} (Assignee: {t.assignee or 'Unassigned'}, Priority: {t.priority or 'Medium'})" for t in analysis_response.tasks])}

Decisions:
{chr(10).join([f"- {d.decision}" for d in analysis_response.decisions])}

Transcript:
{request.transcript}
"""
                
                # Store in ChromaDB
                meetings_collection.add(
                    documents=[document_text],
                    ids=[meeting_id],
                    metadatas=[{
                        "meeting_title": request.meeting_title,
                        "date": datetime.now().isoformat(),
                        "num_tasks": len(analysis_response.tasks),
                        "num_decisions": len(analysis_response.decisions)
                    }]
                )
            except Exception as e:
                print(f"Warning: Failed to store in ChromaDB: {e}")
        
        return analysis_response
        
    except ValueError as e:
        raise HTTPException(status_code=500, detail=f"JSON parsing error: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error analyzing transcript: {str(e)}")

class QueryRequest(BaseModel):
    query: str

class QueryResponse(BaseModel):
    answer: str
    relevant_meetings: List[dict]

@app.post("/api/query", response_model=QueryResponse)
async def query_meetings(request: QueryRequest):
    """
    RAG-based query endpoint for asking questions about past meetings
    Example: "Who is responsible for Task X?" or "What was decided in the Q1 meeting?"
    """
    try:
        if not meetings_collection:
            raise HTTPException(status_code=503, detail="ChromaDB not available")
        
        # Query ChromaDB for relevant meetings
        results = meetings_collection.query(
            query_texts=[request.query],
            n_results=3
        )
        
        if not results['documents'] or not results['documents'][0]:
            return QueryResponse(
                answer="No relevant meetings found for your query.",
                relevant_meetings=[]
            )
        
        # Prepare context from retrieved documents
        context = "\n\n---\n\n".join(results['documents'][0])
        
        # Use LangChain to generate answer
        answer_prompt = ChatPromptTemplate.from_template("""
Based on the following meeting information, answer the user's question concisely and accurately.

Meeting Information:
{context}

User Question: {query}

Provide a clear, direct answer. If the information is not available, say so.""")
        
        chain = answer_prompt | llm
        response = chain.invoke({
            "context": context,
            "query": request.query
        })
        
        # Prepare relevant meetings info
        relevant_meetings = []
        for i, (doc, metadata) in enumerate(zip(results['documents'][0], results['metadatas'][0])):
            relevant_meetings.append({
                "meeting_title": metadata.get("meeting_title", f"Meeting {i+1}"),
                "date": metadata.get("date", "Unknown"),
                "num_tasks": metadata.get("num_tasks", 0),
                "num_decisions": metadata.get("num_decisions", 0)
            })
        
        return QueryResponse(
            answer=response.content,
            relevant_meetings=relevant_meetings
        )
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error querying meetings: {str(e)}")

class ExportToSheetsRequest(BaseModel):
    meeting_title: str
    summary: str
    tasks: List[Task]
    decisions: List[Decision]
    spreadsheet_id: Optional[str] = None

@app.post("/api/export/sheets")
async def export_to_google_sheets(request: ExportToSheetsRequest):
    """
    Export meeting analysis to Google Sheets
    Note: Requires Google Sheets API credentials (credentials.json)
    """
    try:
        import gspread
        from oauth2client.service_account import ServiceAccountCredentials
        
        # Check if credentials file exists
        if not os.path.exists("credentials.json"):
            return {
                "success": False,
                "message": "Google Sheets credentials.json not found. Please add it to the backend directory."
            }
        
        # Setup credentials
        scope = ['https://spreadsheets.google.com/feeds',
                'https://www.googleapis.com/auth/drive']
        creds = ServiceAccountCredentials.from_json_keyfile_name('credentials.json', scope)
        client = gspread.authorize(creds)
        
        # Open or create spreadsheet
        if request.spreadsheet_id:
            sheet = client.open_by_key(request.spreadsheet_id).sheet1
        else:
            spreadsheet = client.create(f"Meeting Analysis - {request.meeting_title}")
            sheet = spreadsheet.sheet1
            spreadsheet_id = spreadsheet.id
        
        # Write headers and data
        sheet.clear()
        sheet.append_row(["Meeting Title", request.meeting_title])
        sheet.append_row(["Date", datetime.now().strftime('%Y-%m-%d %H:%M')])
        sheet.append_row(["Summary", request.summary])
        sheet.append_row([])  # Empty row
        
        # Tasks section
        sheet.append_row(["TASKS"])
        sheet.append_row(["Task", "Assignee", "Deadline", "Priority", "Status"])
        for task in request.tasks:
            sheet.append_row([
                task.task,
                task.assignee or "Unassigned",
                task.deadline or "No deadline",
                task.priority or "Medium",
                "Pending"
            ])
        
        sheet.append_row([])  # Empty row
        
        # Decisions section
        sheet.append_row(["DECISIONS"])
        sheet.append_row(["Decision", "Rationale"])
        for decision in request.decisions:
            sheet.append_row([
                decision.decision,
                decision.rationale or "N/A"
            ])
        
        return {
            "success": True,
            "message": "Successfully exported to Google Sheets",
            "spreadsheet_url": f"https://docs.google.com/spreadsheets/d/{spreadsheet_id if not request.spreadsheet_id else request.spreadsheet_id}"
        }
        
    except Exception as e:
        return {
            "success": False,
            "message": f"Error exporting to Google Sheets: {str(e)}"
        }

@app.post("/api/export/excel")
async def export_to_excel(request: ExportToSheetsRequest):
    """
    Export meeting analysis to Excel file
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Meeting Analysis"
        
        # Add meeting info
        ws['A1'] = "Meeting Title"
        ws['B1'] = request.meeting_title
        ws['A1'].font = Font(bold=True)
        
        ws['A2'] = "Date"
        ws['B2'] = datetime.now().strftime('%Y-%m-%d %H:%M')
        ws['A2'].font = Font(bold=True)
        
        ws['A3'] = "Summary"
        ws['B3'] = request.summary
        ws['A3'].font = Font(bold=True)
        ws['B3'].alignment = Alignment(wrap_text=True)
        
        # Tasks section
        row = 5
        ws[f'A{row}'] = "TASKS"
        ws[f'A{row}'].font = Font(bold=True, size=14)
        row += 1
        
        headers = ["Task", "Assignee", "Deadline", "Priority", "Status"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row, col, header)
            cell.font = Font(bold=True)
        row += 1
        
        for task in request.tasks:
            ws.cell(row, 1, task.task)
            ws.cell(row, 2, task.assignee or "Unassigned")
            ws.cell(row, 3, task.deadline or "No deadline")
            ws.cell(row, 4, task.priority or "Medium")
            ws.cell(row, 5, "Pending")
            row += 1
        
        # Decisions section
        row += 2
        ws[f'A{row}'] = "DECISIONS"
        ws[f'A{row}'].font = Font(bold=True, size=14)
        row += 1
        
        ws.cell(row, 1, "Decision").font = Font(bold=True)
        ws.cell(row, 2, "Rationale").font = Font(bold=True)
        row += 1
        
        for decision in request.decisions:
            ws.cell(row, 1, decision.decision)
            ws.cell(row, 2, decision.rationale or "N/A")
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        
        # Save file
        filename = f"meeting_{request.meeting_title.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        filepath = os.path.join("exports", filename)
        os.makedirs("exports", exist_ok=True)
        wb.save(filepath)
        
        return {
            "success": True,
            "message": "Successfully exported to Excel",
            "filename": filename,
            "filepath": filepath
        }
        
    except Exception as e:
        return {
            "success": False,
            "message": f"Error exporting to Excel: {str(e)}"
        }

@app.get("/health")
def health_check():
    """Health check endpoint"""
    return {
        "status": "healthy",
        "gemini_configured": bool(os.getenv("GEMINI_API_KEY")),
        "chromadb_available": meetings_collection is not None,
        "google_sheets_configured": os.path.exists("credentials.json")
    }

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
